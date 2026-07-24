# -*- coding: utf-8 -*-
"""Fill the 10 UI keys that only had English, then re-export catalog.json."""
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\시나브로-vibe")
XLSX = ROOT / "docs" / "sinavro-i18n.xlsx"
DATA = ROOT / "scripts" / "_missing_ui_fill.json"
EXPORT = ROOT / "scripts" / "export_i18n_json.py"

LANGS = [
    "ja", "zh-CN", "zh-TW", "vi", "th", "id", "mn", "ru", "uz", "kk", "ky",
    "ne", "my", "km", "hi", "bn", "ar", "es", "fr", "de", "sw", "ha", "fil",
]


def main() -> None:
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    en = payload["en"]
    fills = payload["fills"]

    for key, translations in fills.items():
        miss = [code for code in LANGS if not (translations.get(code) or "").strip()]
        if miss:
            raise SystemExit(f"Missing locales for {key}: {miss}")

    wb = load_workbook(XLSX)
    st = wb["strings"]
    headers = [c.value for c in st[1]]
    hmap = {h: i + 1 for i, h in enumerate(headers) if h}

    existing = {}
    for r in range(2, st.max_row + 1):
        key = st.cell(r, 1).value
        if key:
            existing[key] = r

    for key, translations in fills.items():
        if key in existing:
            row = existing[key]
        else:
            row = st.max_row + 1
            st.cell(row, hmap["key"]).value = key
            st.cell(row, hmap["category"]).value = "ui"
            st.cell(row, hmap["en"]).value = en[key]
            if "status" in hmap:
                st.cell(row, hmap["status"]).value = "filled"
            existing[key] = row
        for code in LANGS:
            st.cell(row, hmap[code]).value = translations[code]

    wb.save(XLSX)
    print(f"Updated {XLSX}")

    # re-export catalog
    import runpy

    runpy.run_path(str(EXPORT), run_name="__main__")


if __name__ == "__main__":
    main()
