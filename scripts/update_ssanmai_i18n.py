# -*- coding: utf-8 -*-
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\시나브로-vibe")
DATA = json.loads((ROOT / "scripts" / "_ssanmai_fill.json").read_text(encoding="utf-8"))
KEYS = {
    "era": "card.p1-08.era",
    "meaning": "card.p1-08.meaning",
    "scenario": "card.p1-08.scenario",
}
LANGS = [
    "ja", "zh-CN", "zh-TW", "vi", "th", "id", "mn", "ru", "uz", "kk", "ky",
    "ne", "my", "km", "hi", "bn", "ar", "es", "fr", "de", "sw", "ha", "fil",
]


def update_catalog() -> None:
    path = ROOT / "app" / "i18n" / "catalog.json"
    catalog = json.loads(path.read_text(encoding="utf-8"))
    for field, key in KEYS.items():
        entry = {"en": DATA[field]["en"]}
        for code in LANGS:
            entry[code] = DATA[field][code]
        catalog["strings"][key] = entry
    path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("catalog updated")


def update_xlsx() -> None:
    xlsx = ROOT / "docs" / "sinavro-i18n.xlsx"
    wb = load_workbook(xlsx)
    st = wb["strings"]
    headers = [c.value for c in st[1]]
    hmap = {h: i + 1 for i, h in enumerate(headers) if h}
    rows = {}
    for r in range(2, st.max_row + 1):
        key = st.cell(r, 1).value
        if key in KEYS.values():
            rows[key] = r
    for field, key in KEYS.items():
        row = rows.get(key)
        if row is None:
            raise SystemExit(f"Missing row {key}")
        st.cell(row, hmap["en"]).value = DATA[field]["en"]
        for code in LANGS:
            st.cell(row, hmap[code]).value = DATA[field][code]
    # keep_ko sheet if it mentions 느좋
    if "keep_ko" in wb.sheetnames:
        ks = wb["keep_ko"]
        for r in range(1, ks.max_row + 1):
            for c in range(1, ks.max_column + 1):
                val = ks.cell(r, c).value
                if isinstance(val, str) and "느좋" in val:
                    ks.cell(r, c).value = val.replace("느좋", "싼마이")
    wb.save(xlsx)
    print("xlsx updated")


def main() -> None:
    # validate no bad chars
    for field, block in DATA.items():
        for code, text in block.items():
            if "\ufffd" in text:
                raise SystemExit(f"Bad char in {field}.{code}")
            if code != "en" and code not in LANGS:
                raise SystemExit(f"Unexpected locale {code}")
        miss = [c for c in LANGS if c not in block]
        if miss:
            raise SystemExit(f"Missing locales in {field}: {miss}")
    update_catalog()
    update_xlsx()


if __name__ == "__main__":
    main()
