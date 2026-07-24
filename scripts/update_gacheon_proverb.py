# -*- coding: utf-8 -*-
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\시나브로-vibe")
DATA = json.loads((ROOT / "scripts" / "_gacheon_fill.json").read_text(encoding="utf-8"))
KO = DATA["ko"]
FILLS = DATA["fills"]


def update_xlsx() -> None:
    xlsx = ROOT / "docs" / "sinavro-i18n.xlsx"
    wb = load_workbook(xlsx)
    pm = wb["proverb_map"]
    headers = [c.value for c in pm[1]]
    hmap = {h: i + 1 for i, h in enumerate(headers) if h}

    row = None
    for r in range(2, pm.max_row + 1):
        if pm.cell(r, 1).value == KO:
            row = r
            break
    if row is None:
        raise SystemExit(f"Row not found for {KO}")

    if "en" in hmap:
        pm.cell(row, hmap["en"]).value = DATA["en"]
    if "en_type" in hmap:
        pm.cell(row, hmap["en_type"]).value = DATA["en_type"]

    for code, item in FILLS.items():
        if code not in hmap:
            raise SystemExit(f"Missing column {code}")
        pm.cell(row, hmap[code]).value = item["text"]
        type_col = f"{code}_type"
        if type_col in hmap:
            pm.cell(row, hmap[type_col]).value = item["type"]

    wb.save(xlsx)
    print(f"xlsx updated row {row}")


def update_catalog() -> None:
    path = ROOT / "app" / "i18n" / "catalog.json"
    catalog = json.loads(path.read_text(encoding="utf-8"))
    entry = catalog["proverbs"][KO]
    entry["en"] = DATA["en"]
    types = entry.setdefault("types", {})
    types["en"] = DATA["en_type"]
    for code, item in FILLS.items():
        entry[code] = item["text"]
        types[code] = item["type"]
    path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("catalog updated")


def main() -> None:
    update_catalog()
    try:
        update_xlsx()
    except PermissionError:
        print("WARN: xlsx locked; catalog updated only. Close Excel and re-run.")


if __name__ == "__main__":
    main()
