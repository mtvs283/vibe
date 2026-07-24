# -*- coding: utf-8 -*-
import json
from openpyxl import load_workbook
from pathlib import Path

wb = load_workbook(r"c:\시나브로-vibe\docs\sinavro-i18n.xlsx")
B = ["ne", "my", "km", "hi", "bn", "ar", "es", "fr", "de", "sw", "ha", "fil"]
st = wb["strings"]
h = [c.value for c in st[1]]
hm = {x: i + 1 for i, x in enumerate(h)}
pm = wb["proverb_map"]
ph = [c.value for c in pm[1]]
pmh = {x: i + 1 for i, x in enumerate(ph)}

issues = []
for lang in B:
    for r in range(2, st.max_row + 1):
        en = st.cell(r, 3).value or ""
        val = st.cell(r, hm[lang]).value or ""
        key = st.cell(r, 1).value
        if not val:
            issues.append({"lang": lang, "key": key, "issue": "EMPTY"})
            continue
        if "{n}" in en and "{n}" not in val:
            issues.append({"lang": lang, "key": key, "issue": "MISSING_{n}", "val": val})

for lang in B:
    for r in range(2, pm.max_row + 1):
        ko = pm.cell(r, 1).value
        en = pm.cell(r, 2).value
        text = pm.cell(r, pmh[lang]).value or ""
        typ = pm.cell(r, pmh[f"{lang}_type"]).value or ""
        note = pm.cell(r, pmh[f"{lang}_note"]).value or ""
        if not text:
            issues.append({"lang": lang, "key": ko, "issue": "PROVERB_EMPTY"})
            continue
        if not typ:
            issues.append({"lang": lang, "key": ko, "issue": "MISSING_TYPE", "val": text})
        if text.strip() == (en or "").strip() and lang in ("ne", "my", "hi", "bn", "ar", "ha"):
            issues.append({"lang": lang, "key": ko, "issue": "PROVERB_SAME_AS_EN", "val": text})
        if typ == "paraphrase" and not note:
            issues.append({"lang": lang, "key": ko, "issue": "PARAPHRASE_NO_NOTE", "val": text})

samples = {}
targets = [
    "제 무덤 내가 판다",
    "똥 묻은 개가 겨 묻은 개 나무란다",
    "설마가 사람 잡는다",
    "호랑이 담배 피던 시절",
    "참새가 방앗간을 그냥 지나치랴",
    "벼는 익을수록 고개를 숙인다",
    "개똥도 약에 쓰려면 없다",
    "남의 떡이 커 보인다",
]
for r in range(2, pm.max_row + 1):
    ko = pm.cell(r, 1).value
    if ko in targets:
        samples[ko] = {"en": pm.cell(r, 2).value}
        for lang in B:
            samples[ko][lang] = {
                "t": pm.cell(r, pmh[lang]).value,
                "type": pm.cell(r, pmh[f"{lang}_type"]).value,
                "note": pm.cell(r, pmh[f"{lang}_note"]).value,
            }

ui = {}
for key in [
    "ui.brand_slogan",
    "ui.result.misses_line",
    "ui.find_proverb_hint_p2",
    "card.p1-01.meaning",
    "card.p1-04.meaning",
]:
    row = next(r for r in range(2, st.max_row + 1) if st.cell(r, 1).value == key)
    ui[key] = {"en": st.cell(row, 3).value}
    for lang in B:
        ui[key][lang] = st.cell(row, hm[lang]).value

counts = {}
for i in issues:
    counts[i["lang"]] = counts.get(i["lang"], 0) + 1

# qualitative flags I add after reading samples
qual = []

# Check Back in my day / Famous last words style
for ko in targets:
    s = samples[ko]
    # es/fr/de should not be literal Korean calques for tiger smoking unless marked paraphrase
    for lang in ("es", "fr", "de"):
        t = (s[lang]["t"] or "").lower()
        if ko == "호랑이 담배 피던 시절":
            # good if has "tiempo"/"jour"/"Tag" old days; bad if only tiger smoke without note
            if "tiger" in t or "tigre" in t or "tiger" in (s[lang]["t"] or "").lower() or "tigre" in (s[lang]["t"] or "").lower():
                if s[lang]["type"] != "paraphrase" and not (s[lang]["note"] or ""):
                    qual.append(f"{lang} | {ko} | tiger-literal without note?")
        if ko == "설마가 사람 잡는다":
            # famous last words equivalents
            pass

out = {
    "issue_counts": counts,
    "issues": issues,
    "qual_flags": qual,
    "samples_proverb": samples,
    "samples_ui": ui,
}
Path(r"c:\시나브로-vibe\docs\_qa_B_review.json").write_text(
    json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8"
)

# readable summary
lines = ["# B-column review (by Cursor/A)", ""]
lines.append("## Completeness")
for lang in B:
    lines.append(f"- {lang}: issues={counts.get(lang,0)}")
lines.append("")
lines.append("## Critical")
lines.append("- km: strings 123 empty, proverbs 42 empty (NOT DONE)")
lines.append("")
lines.append("## Proverb spot-check")
for ko in targets:
    lines.append(f"### {ko}")
    lines.append(f"en: {samples[ko]['en']}")
    for lang in ["es", "fr", "de", "ar", "hi", "fil", "sw", "ha", "ne", "my", "bn"]:
        x = samples[ko][lang]
        lines.append(f"- {lang} [{x['type']}] {x['t']} | note={x['note']}")
    lines.append("")
lines.append("## UI spot-check")
for key, val in ui.items():
    lines.append(f"### {key}")
    lines.append(f"en: {val['en']}")
    for lang in ["es", "fr", "de", "ar", "hi", "fil"]:
        lines.append(f"- {lang}: {val[lang]}")
    lines.append("")

Path(r"c:\시나브로-vibe\docs\B_REVIEW.md").write_text("\n".join(lines), encoding="utf-8")
print("issues", counts)
print("qual", qual)
print("km empty?", counts.get("km"))
