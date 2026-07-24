# -*- coding: utf-8 -*-
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(r"c:\시나브로-vibe\docs\sinavro-i18n.xlsx")
st = wb["strings"]
h = [c.value for c in st[1]]
hm = {x: i + 1 for i, x in enumerate(h)}
pm = wb["proverb_map"]
ph = [c.value for c in pm[1]]
pmh = {x: i + 1 for i, x in enumerate(ph)}

lines = []
for lang in ("sw", "ha"):
    lines.append(f"## {lang} paraphrase only")
    for r in range(2, pm.max_row + 1):
        ko = pm.cell(r, 1).value
        en = pm.cell(r, 2).value
        text = pm.cell(r, pmh[lang]).value
        typ = pm.cell(r, pmh[f"{lang}_type"]).value
        note = pm.cell(r, pmh[f"{lang}_note"]).value
        if (typ or "").lower() == "paraphrase":
            lines.append(f"- {ko}")
            lines.append(f"  EN: {en}")
            lines.append(f"  {lang}: {text} | note={note}")
    lines.append("")
    lines.append(f"## {lang} all proverbs")
    for r in range(2, pm.max_row + 1):
        ko = pm.cell(r, 1).value
        en = pm.cell(r, 2).value
        text = pm.cell(r, pmh[lang]).value
        typ = pm.cell(r, pmh[f"{lang}_type"]).value
        note = pm.cell(r, pmh[f"{lang}_note"]).value
        lines.append(f"- [{typ}] {ko}")
        lines.append(f"  EN: {en}")
        lines.append(f"  XX: {text} | {note}")
    lines.append("")

lines.append("## UI + card fields")
for r in range(2, st.max_row + 1):
    key = st.cell(r, 1).value
    cat = st.cell(r, 2).value
    if cat not in ("ui", "card"):
        continue
    en = st.cell(r, 3).value
    lines.append(f"KEY {key}")
    lines.append(f" en: {en}")
    lines.append(f" sw: {st.cell(r, hm['sw']).value}")
    lines.append(f" ha: {st.cell(r, hm['ha']).value}")

Path(r"c:\시나브로-vibe\docs\_qa_sw_ha_detail.txt").write_text("\n".join(lines), encoding="utf-8")
print("done", len(lines))
