# -*- coding: utf-8 -*-
import sys
from openpyxl import load_workbook
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(r"c:\시나브로-vibe\docs\sinavro-i18n.xlsx")
st = wb["strings"]
h = [c.value for c in st[1]]
hm = {x: i + 1 for i, x in enumerate(h)}
pm = wb["proverb_map"]
ph = [c.value for c in pm[1]]
pmh = {x: i + 1 for i, x in enumerate(ph)}
B = ["ne", "my", "km", "hi", "bn", "ar", "es", "fr", "de", "sw", "ha", "fil"]

lines = []
lines.append("=== paraphrase rates ===")
for lang in B:
    para = loose = 0
    for r in range(2, pm.max_row + 1):
        typ = pm.cell(r, pmh[f"{lang}_type"]).value or ""
        note = pm.cell(r, pmh[f"{lang}_note"]).value or ""
        if typ == "paraphrase":
            para += 1
        if "loose" in str(note).lower():
            loose += 1
    lines.append(f"{lang}: paraphrase={para}/42 loose_note={loose}/42")

lines.append("=== placeholder failures ===")
fails = 0
for r in range(2, st.max_row + 1):
    en = st.cell(r, 3).value or ""
    if "{n}" not in en:
        continue
    key = st.cell(r, 1).value
    for lang in B:
        val = st.cell(r, hm[lang]).value or ""
        if "{n}" not in val:
            lines.append(f"FAIL {lang} {key} :: {val}")
            fails += 1
if fails == 0:
    lines.append("none")

lines.append("=== UI samples ===")
for key in [
    "ui.brand_slogan",
    "ui.hero_copy",
    "ui.correct",
    "ui.wrong",
    "ui.try_again",
    "ui.result.misses_line",
    "card.p1-01.meaning",
    "card.p1-04.scenario",
]:
    row = next(r for r in range(2, st.max_row + 1) if st.cell(r, 1).value == key)
    lines.append(f"KEY {key}")
    lines.append(f" en: {st.cell(row, 3).value}")
    for lang in ["es", "fr", "de", "ar", "hi", "fil", "km", "sw", "ha", "ne"]:
        lines.append(f" {lang}: {st.cell(row, hm[lang]).value}")

lines.append("=== proverb samples ===")
for ko in [
    "제 무덤 내가 판다",
    "남의 떡이 커 보인다",
    "설마가 사람 잡는다",
    "호랑이 담배 피던 시절",
    "되로 주고 말로 받는다",
    "서당 개 삼 년이면 풍월을 읊는다",
    "개똥도 약에 쓰려면 없다",
    "똥 묻은 개가 겨 묻은 개 나무란다",
]:
    row = next(r for r in range(2, pm.max_row + 1) if pm.cell(r, 1).value == ko)
    lines.append(f"KO {ko} | {pm.cell(row, 2).value}")
    for lang in ["es", "fr", "de", "ar", "hi", "fil", "km", "ne", "my", "sw", "ha", "bn"]:
        lines.append(
            f" {lang}[{pm.cell(row, pmh[f'{lang}_type']).value}]: {pm.cell(row, pmh[lang]).value} | {pm.cell(row, pmh[f'{lang}_note']).value}"
        )

Path = __import__("pathlib").Path
Path(r"c:\시나브로-vibe\docs\_qa_B_samples.txt").write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines[:80]))
print("... wrote full to _qa_B_samples.txt")
