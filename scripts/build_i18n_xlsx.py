# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import os
import re

ROOT = r"c:\시나브로-vibe"
SRC = os.path.join(ROOT, "app", "page.tsx")
OUT = os.path.join(ROOT, "docs", "sinavro-i18n.xlsx")

placeholders = [
    (1, "en", "영어", "source", ""),
    (2, "ja", "일본어", "A", ""),
    (3, "zh-CN", "중국어(간체)", "A", ""),
    (4, "zh-TW", "중국어(번체)", "A", ""),
    (5, "vi", "베트남어", "A", ""),
    (6, "th", "태국어", "A", ""),
    (7, "id", "인도네시아어", "A", ""),
    (8, "ms", "말레이어", "A", ""),
    (9, "tl", "타갈로그", "A", ""),
    (10, "es", "스페인어", "A", ""),
    (11, "pt", "포르투갈어", "A", ""),
    (12, "fr", "프랑스어", "A", ""),
    (13, "de", "독일어", "B", ""),
    (14, "ru", "러시아어", "B", ""),
    (15, "uk", "우크라이나어", "B", ""),
    (16, "pl", "폴란드어", "B", ""),
    (17, "tr", "터키어", "B", ""),
    (18, "ar", "아랍어", "B", ""),
    (19, "hi", "힌디어", "B", ""),
    (20, "bn", "벵골어", "B", ""),
    (21, "uz", "우즈벡어", "B", ""),
    (22, "mn", "몽골어", "B", ""),
    (23, "kk", "카자흐어", "B", ""),
    (24, "sw", "스와힐리어", "B", ""),
]
lang_codes = [p[1] for p in placeholders if p[1] != "en"]

src = open(SRC, encoding="utf-8").read()

wb = Workbook()

# README
ws = wb.active
ws.title = "README"
ws["A1"] = "시나브로 전체 로컬라이즈 (한글 제외)"
ws["A1"].font = Font(bold=True, size=14)
readme = """
대상
- 번역: UI / meaning / scenario / era / format / 속담 대응 표현 / (선택) alt
- 유지: 한국어 신조어·리믹스, 한국어 속담 원문, 로마자, 기관명 한글

시트
1) locales — 언어 24개. en=소스. owner A=나, B=클로드
2) strings — UI·카드 설명 문장 전부
3) proverb_map — 한국어 속담 → 각 언어 속담/관용구
4) keep_ko — 번역 금지 목록

strings 규칙
- key / category / en 수정 금지
- 담당 언어 열만 채움
- 한 셀 = 한 문장

proverb_map 규칙
- 직역보다 그 언어의 속담/관용구 우선
- 없으면 paraphrase + _note에 loose
"""
for i, line in enumerate(readme.strip().split("\n"), start=2):
    ws[f"A{i}"] = line
ws.column_dimensions["A"].width = 100

# locales
loc = wb.create_sheet("locales")
for c, h in enumerate(["order", "code", "name_ko", "owner", "done"], 1):
    cell = loc.cell(1, c, h)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D9EAD3")
for r, row in enumerate(placeholders, 2):
    for c, v in enumerate(row, 1):
        loc.cell(r, c, v)
for col, w in zip("ABCDE", [8, 12, 22, 10, 8]):
    loc.column_dimensions[col].width = w

strings = []

ui = [
    ("ui.edition_pill", "ui", "ENGLISH EDITION"),
    ("ui.sinavro_presents", "ui", "SINAVRO PRESENTS"),
    ("ui.product_descriptor", "ui", "KOREAN PROVERBS & SLANG CARD GAME"),
    ("ui.tap_to_continue", "ui", "TAP TO CONTINUE →"),
    ("ui.brand_slogan", "ui", "Old Wisdom, New Vibes."),
    (
        "ui.hero_copy",
        "ui",
        "Match Korean internet slang with timeless proverbs. Then remix the old wisdom in the language of memes.",
    ),
    ("ui.tag.korean_slang", "ui", "#KoreanSlang"),
    ("ui.tag.proverbs", "ui", "#Proverbs"),
    ("ui.tag.world_wisdom", "ui", "#WorldWisdom"),
    ("ui.part01_title", "ui", "Catch the Vibe"),
    (
        "ui.part01_body",
        "ui",
        "Discover real Korean slang, read the situation, and find its closest proverb cousin.",
    ),
    ("ui.part01_li1", "ui", "8 real expressions"),
    ("ui.part01_li2", "ui", "Two-panel visual cards"),
    ("ui.part01_li3", "ui", "English proverb matches"),
    ("ui.start_part1", "ui", "START PART 1 →"),
    ("ui.part02_title", "ui", "Proverb Remix"),
    (
        "ui.part02_body",
        "ui",
        "Decode Korean proverbs rewritten as memes, community posts, and modern punchlines.",
    ),
    ("ui.part02_li1", "ui", "8 creative remixes"),
    ("ui.part02_li2", "ui", "Two-panel visual cards"),
    ("ui.part02_li3", "ui", "World-wisdom discussion"),
    ("ui.start_part2", "ui", "START PART 2 →"),
    ("ui.game_title.part1", "ui", "CATCH THE VIBE"),
    ("ui.game_title.part2", "ui", "PROVERB REMIX"),
    ("ui.misses", "ui", "MISSES"),
    ("ui.part_switch_1", "ui", "← PART 1"),
    ("ui.part_switch_2", "ui", "PART 2 →"),
    ("ui.part_complete", "ui", "PART {n} COMPLETE"),
    ("ui.result.immaculate", "ui", "Your vibe is immaculate."),
    ("ui.result.close", "ui", "Close. The wisdom is landing."),
    ("ui.result.loading", "ui", "The wisdom is loading."),
    ("ui.result.zero_misses", "ui", "Zero misses. You read every vibe on the first try."),
    (
        "ui.result.misses_line",
        "ui",
        "You missed {n} time(s). Lower is better—try again for a cleaner run.",
    ),
    (
        "ui.result.after_part1",
        "ui",
        "You have decoded the slang. Now use those new vibes to remix old Korean wisdom.",
    ),
    (
        "ui.result.after_part2",
        "ui",
        "You have travelled from old proverbs to internet memes—and found the human truth underneath both.",
    ),
    ("ui.around_the_world", "ui", "AROUND THE WORLD"),
    (
        "ui.world_question",
        "ui",
        "Does your language have a proverb that matched one of today’s cards?",
    ),
    ("ui.continue_part2", "ui", "CONTINUE TO PART 2 →"),
    ("ui.go_part1", "ui", "← GO TO PART 1"),
    ("ui.play_again", "ui", "PLAY AGAIN"),
    ("ui.back_home", "ui", "BACK TO HOME"),
    ("ui.old_wisdom_new_vibes", "ui", "OLD WISDOM, NEW VIBES"),
    ("ui.correct", "ui", "CORRECT"),
    ("ui.wrong", "ui", "WRONG"),
    ("ui.try_again", "ui", "TRY AGAIN"),
    ("ui.next", "ui", "NEXT"),
    ("ui.decode_it", "ui", "DECODE IT"),
    ("ui.find_proverb_q", "ui", "What is the original proverb?"),
    ("ui.find_proverb_hint_p1", "ui", "Find the Korean proverb that matches this vibe."),
    (
        "ui.find_proverb_hint_p2",
        "ui",
        "Find the original Korean proverb hiding inside this modern remix.",
    ),
    ("ui.image_slot", "ui", "Image slot"),
    ("ui.image_slot_hint", "ui", "Drop Gemini art here later."),
    ("ui.part1_cleared", "ui", "PART 1 CLEARED ✓"),
]
strings.extend(ui)

part1_cards = [
    (
        "p1-01",
        "2020s internet slang",
        "I brought this disaster upon myself.",
        "You ignored every warning. Now you are standing in the mess you made with your own hands.",
    ),
    (
        "p1-02",
        "2020s lifestyle slang",
        "A disciplined, productive, admirable life.",
        "Wake up early, exercise, study Korean, prepare lunch, and still arrive at work on time.",
    ),
    (
        "p1-03",
        "2022 sports & gaming meme",
        "What matters is an unbreakable spirit.",
        "Your team is far behind, but nobody gives up before the final whistle.",
    ),
    (
        "p1-04",
        "2010s–2020s public discourse",
        "One rule for me, another for you.",
        "He complains when others are five minutes late, but expects everyone to wait when he is late.",
    ),
    (
        "p1-05",
        "2020s reaction meme",
        "This might actually be better!",
        "You miss the crowded bus in the rain—then the next one gives you a luxury empty seat.",
    ),
    (
        "p1-06",
        "late 2010s–2020s slang",
        "Handle it well, neatly, and with good sense—without being told every detail.",
        "You give one short instruction. Your coworker understands the whole situation and delivers perfectly.",
    ),
    (
        "p1-07",
        "2000s comparison culture",
        "The impossibly perfect ‘friend’s son’ your mother compares you with.",
        "Your family praises someone else’s grades, job, manners, cooking, and fitness—all during dinner.",
    ),
    (
        "p1-08",
        "2020s consumer slang",
        "Cheap-looking, low-quality, no-good vibes.",
        "The rice cakes look bargain-cheap in the case—then one bite proves why.",
    ),
]
for key, era, meaning, scenario in part1_cards:
    strings.append((f"card.{key}.era", "card", era))
    strings.append((f"card.{key}.meaning", "card", meaning))
    strings.append((f"card.{key}.scenario", "card", scenario))

p2 = re.findall(
    r'remix:\s*"([^"]+)"[\s\S]*?format:\s*"([^"]+)"[\s\S]*?answer:\s*"([^"]+)"[\s\S]*?meaning:\s*"([^"]+)"',
    src,
)
for i, (_remix, format_, _answer, meaning) in enumerate(p2, 1):
    key = f"p2-{i:02d}"
    strings.append((f"card.{key}.format", "card", format_))
    strings.append((f"card.{key}.meaning", "card", meaning))

alts = re.findall(r'alt:\s*"([^"]+)"', src)
for i, alt in enumerate(alts, 1):
    strings.append((f"alt.{i:03d}", "alt", alt))

opt_pairs = re.findall(r'\{\s*ko:\s*"([^"]+)",\s*en:\s*"([^"]+)"\s*\}', src)
proverb_map = {}
for ko, en in opt_pairs:
    proverb_map.setdefault(ko, en)

st = wb.create_sheet("strings")
headers = ["key", "category", "en"] + lang_codes + ["status"]
fill = PatternFill("solid", fgColor="FFF2CC")
for c, h in enumerate(headers, 1):
    cell = st.cell(1, c, h)
    cell.font = Font(bold=True)
    cell.fill = fill

seen = set()
rows = []
for key, cat, en in strings:
    if key in seen:
        continue
    seen.add(key)
    rows.append((key, cat, en))

for r, (key, cat, en) in enumerate(rows, 2):
    st.cell(r, 1, key)
    st.cell(r, 2, cat)
    st.cell(r, 3, en)
    st.cell(r, len(headers), "empty")

st.column_dimensions["A"].width = 36
st.column_dimensions["B"].width = 10
st.column_dimensions["C"].width = 70
for c in range(4, len(headers)):
    st.column_dimensions[get_column_letter(c)].width = 28
st.freeze_panes = "D2"

pm = wb.create_sheet("proverb_map")
pm_headers = ["ko_proverb", "en_match", "en_type"]
for code in lang_codes:
    pm_headers += [code, f"{code}_type", f"{code}_note"]
for c, h in enumerate(pm_headers, 1):
    cell = pm.cell(1, c, h)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D0E0FF")
for r, (ko, en) in enumerate(sorted(proverb_map.items(), key=lambda x: x[0]), 2):
    pm.cell(r, 1, ko)
    pm.cell(r, 2, en)
    pm.cell(r, 3, "")
pm.column_dimensions["A"].width = 40
pm.column_dimensions["B"].width = 45
pm.column_dimensions["C"].width = 12
for c in range(4, 16):
    pm.column_dimensions[get_column_letter(c)].width = 26
pm.freeze_panes = "D2"

kk = wb.create_sheet("keep_ko")
kk["A1"] = "do_not_translate"
kk["B1"] = "example"
kk["A1"].font = Font(bold=True)
for r, row in enumerate(
    [
        ("신조어/리믹스 원문", "스불재, 라떼는 말이야"),
        ("한국어 속담 원문", "제 무덤 내가 판다"),
        ("로마자", "seu-bul-jae"),
        ("기관명", "한국어교육AI연구개발원"),
        ("브랜드명", "SINAVRO"),
    ],
    2,
):
    kk.cell(r, 1, row[0])
    kk.cell(r, 2, row[1])
kk.column_dimensions["A"].width = 28
kk.column_dimensions["B"].width = 40

os.makedirs(os.path.join(ROOT, "docs"), exist_ok=True)
wb.save(OUT)
print(OUT)
print("strings", len(rows), "proverbs", len(proverb_map))
