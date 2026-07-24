# -*- coding: utf-8 -*-
"""Export sinavro-i18n.xlsx → app/i18n/catalog.json"""
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\시나브로-vibe")
XLSX = ROOT / "docs" / "sinavro-i18n.xlsx"
OUT = ROOT / "app" / "i18n" / "catalog.json"

wb = load_workbook(XLSX, data_only=True)

# locales
loc = wb["locales"]
locales = []
for r in range(2, loc.max_row + 1):
    code = loc.cell(r, 2).value
    name = loc.cell(r, 3).value
    owner = loc.cell(r, 4).value
    if not code or code == "TBD":
        continue
    locales.append({"code": code, "nameKo": name, "owner": owner})

# ensure en first
locales.sort(key=lambda x: (0 if x["code"] == "en" else 1, x["code"]))

st = wb["strings"]
headers = [c.value for c in st[1]]
hmap = {h: i + 1 for i, h in enumerate(headers) if h}
lang_cols = [h for h in headers if h not in ("key", "category", "en", "status", None)]

strings = {}
for r in range(2, st.max_row + 1):
    key = st.cell(r, 1).value
    if not key:
        continue
    entry = {"en": st.cell(r, 3).value or ""}
    for code in lang_cols:
        val = st.cell(r, hmap[code]).value
        if val is not None and str(val).strip():
            entry[code] = str(val)
    strings[key] = entry

pm = wb["proverb_map"]
ph = [c.value for c in pm[1]]
pmap = {h: i + 1 for i, h in enumerate(ph) if h}
proverbs = {}
for r in range(2, pm.max_row + 1):
    ko = pm.cell(r, 1).value
    if not ko:
        continue
    entry = {
        "en": pm.cell(r, 2).value or "",
        "types": {},
        "notes": {},
    }
    for code in lang_cols:
        if code not in pmap:
            continue
        text = pm.cell(r, pmap[code]).value
        typ = pm.cell(r, pmap.get(f"{code}_type", 0)).value if f"{code}_type" in pmap else None
        note = pm.cell(r, pmap.get(f"{code}_note", 0)).value if f"{code}_note" in pmap else None
        if text is not None and str(text).strip():
            entry[code] = str(text)
        if typ:
            entry["types"][code] = str(typ)
        if note:
            entry["notes"][code] = str(note)
    # clean empty dicts
    if not entry["types"]:
        del entry["types"]
    if not entry["notes"]:
        del entry["notes"]
    proverbs[ko] = entry

catalog = {
    "version": 1,
    "locales": locales,
    "strings": strings,
    "proverbs": proverbs,
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
print("wrote", OUT)
print("locales", len(locales), "strings", len(strings), "proverbs", len(proverbs))
